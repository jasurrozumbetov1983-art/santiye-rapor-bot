import os
import re
from datetime import datetime
from zoneinfo import ZoneInfo

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from telegram import Update, ReplyKeyboardMarkup
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    ContextTypes,
    filters,
)

load_dotenv()

BOT_TOKEN = os.getenv("BOT_TOKEN")
ADMIN_CHAT_ID = os.getenv("ADMIN_CHAT_ID")
TZ = ZoneInfo(os.getenv("TIMEZONE", "Asia/Tashkent"))

REPORT_HOUR = int(os.getenv("REPORT_HOUR", "15"))
REPORT_MINUTE = int(os.getenv("REPORT_MINUTE", "0"))

# Current/live reports file
DATA_FILE = "raporlar.xlsx"

# Historical Telegram export converted to Excel
HISTORY_FILE = "raporlar_tarih_2026_2.xlsx"

SITES = [
    "DATA CENTER", "DMC", "LOT13", "LOT71", "SKP", "STADYUM",
    "BWC", "KÖKSARAY", "PMP", "MOS", "PİRAMİT", "RMC", "TYM", "YHP"
]

ALIASES = {
    "DATA CENTER": "DATA CENTER", "DATA CENTER": "DATA CENTER",
    "STADIUM": "STADYUM", "STADYUM": "STADYUM",
    "PIRAMIT": "PİRAMİT", "PİRAMİT": "PİRAMİT",
    "PIRAMIT TOWER": "PİRAMİT", "PİRAMİT TOWER": "PİRAMİT",
    "BWC": "BWC", "SKP": "SKP", "LOT71": "LOT71", "LOT13": "LOT13",
    "DMC": "DMC", "PMP": "PMP", "MOS": "MOS", "RMC": "RMC",
    "TYM": "TYM", "YHP": "YHP",
    "KOKSARAY": "KÖKSARAY", "KÖKSARAY": "KÖKSARAY",
    "ELLIPSE GARDEN": "ELLIPSE GARDEN",
}

if "ELLIPSE GARDEN" not in SITES:
    SITES.append("ELLIPSE GARDEN")

MENU = ReplyKeyboardMarkup(
    [["📊 Rapor Durumu", "📥 Excel"], ["🔎 Rapor Ara"]],
    resize_keyboard=True
)


def normalize(s):
    return (
        str(s or "").upper()
        .replace("İ", "I")
        .replace("Ş", "S")
        .replace("Ğ", "G")
        .replace("Ü", "U")
        .replace("Ö", "O")
        .replace("Ç", "C")
    )


def detect_site(text):
    n = normalize(text)
    for alias, site in sorted(ALIASES.items(), key=lambda x: len(x[0]), reverse=True):
        if normalize(alias) in n:
            return site
    return None


def detect_total_workers(text):
    n = normalize(text)
    patterns = [
        r"GENEL\s+TOPLAM\s*[:\-]?\s*(\d+)\s*(?:KISI|KİŞİ)?",
        r"GENEL TOPLAM.*?(\d+)",
        r"TOPLAM\s*[:\-]?\s*(\d+)\s*(?:KISI|KİŞİ)",
    ]
    for pattern in patterns:
        m = re.search(pattern, n, re.S)
        if m:
            try:
                return int(m.group(1))
            except ValueError:
                pass
    return None


def ensure_excel():
    if os.path.exists(DATA_FILE):
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Raporlar"
    ws.append(["Tarih", "Saat", "Şantiye", "Gönderen", "İşçi Sayısı", "Rapor Metni", "Mesaj ID"])
    wb.save(DATA_FILE)


def save_report(site, sender, workers, text, msg_id):
    ensure_excel()
    wb = load_workbook(DATA_FILE)
    ws = wb["Raporlar"]
    now = datetime.now(TZ)
    ws.append([
        now.strftime("%d.%m.%Y"),
        now.strftime("%H:%M"),
        site,
        sender,
        workers if workers is not None else "",
        text,
        msg_id,
    ])
    wb.save(DATA_FILE)


def today_reports():
    ensure_excel()
    wb = load_workbook(DATA_FILE, read_only=True)
    ws = wb["Raporlar"]
    today = datetime.now(TZ).strftime("%d.%m.%Y")
    result = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == today and row[2]:
            result[row[2]] = row
    return result


def status_text():
    reports = today_reports()
    sent = [s for s in SITES if s in reports]
    missing = [s for s in SITES if s not in reports]

    t = f"🕒 {datetime.now(TZ).strftime('%H:%M')} Şantiye Rapor Durumu\n\n"
    t += f"✅ Rapor ileten şantiyeler ({len(sent)}):\n"
    t += ", ".join(sent) if sent else "Menü yok"
    t += f"\n\n❌ Rapor iletmeyen şantiyeler ({len(missing)}):\n"
    t += ", ".join(missing) if missing else "Eksik rapor yok"
    t += "\n\n📌 Not: Yapılan işin raporunu vermek, işi yapmak kadar önemlidir."
    return t


def history_rows():

    if not os.path.exists(HISTORY_FILE):

        return []

    wb = load_workbook(HISTORY_FILE, read_only=True, data_only=True)

    rows = []

    date_pattern = re.compile(r"\b(\d{1,2})[./-](\d{1,2})[./-](\d{4})\b")

    for ws in wb.worksheets:

        for excel_row in ws.iter_rows(values_only=True):

            values = [str(v).strip() for v in excel_row if v not in (None, "")]

            if not values:

                continue

            full_text = " | ".join(values)

            # Sanani topish

            m = date_pattern.search(full_text)

            if not m:

                continue

            day, month, year = m.groups()

            try:

                date_obj = datetime(

                    int(year),

                    int(month),

                    int(day)

                )

            except ValueError:

                continue

            date_str = date_obj.strftime("%d.%m.%Y")

            # Şantiye aniqlash

            site = detect_site(full_text)

            # Ishchilar soni

            workers = detect_total_workers(full_text)

            # Agar sayt topilmasa, bu qatorni ham saqlaymiz,

            # chunki keyinchalik matn bo'yicha qidirish mumkin.

            rows.append((

                date_str,

                site or "",

                workers if workers is not None else "",

                "",

                str(ws.title),

                full_text

            ))

    return rows

def search_history(query):

    query = str(query or "").strip()

    if not query:

        return []

    nq = normalize(query)

    rows = history_rows()

    # Sana bo'yicha qidiruv

    dates = re.findall(

        r"\b\d{1,2}[./-]\d{1,2}[./-]\d{4}\b",

        query

    )

    if len(dates) == 1:

        target = dates[0].replace("/", ".").replace("-", ".")

        try:

            target_dt = datetime.strptime(

                target,

                "%d.%m.%Y"

            ).date()

        except ValueError:

            return []

        filtered = []

        for row in rows:

            try:

                row_dt = datetime.strptime(

                    row[0],

                    "%d.%m.%Y"

                ).date()

                if row_dt == target_dt:

                    filtered.append(row)

            except Exception:

                pass

        return filtered

    # Sana oralig'i

    if len(dates) >= 2:

        try:

            start = datetime.strptime(

                dates[0].replace("/", ".").replace("-", "."),

                "%d.%m.%Y"

            ).date()

            end = datetime.strptime(

                dates[1].replace("/", ".").replace("-", "."),

                "%d.%m.%Y"

            ).date()

        except ValueError:

            return []

        if start > end:

            start, end = end, start

        filtered = []

        for row in rows:

            try:

                row_dt = datetime.strptime(

                    row[0],

                    "%d.%m.%Y"

                ).date()

                if start <= row_dt <= end:

                    filtered.append(row)

            except Exception:

                pass

        return filtered

    # Shantiye / matn / ism bo'yicha qidiruv

    filtered = []

    for row in rows:

        searchable = " ".join(

            str(x or "")

            for x in row

        )

        if nq in normalize(searchable):

            filtered.append(row)

    return filtered

def format_history_results(rows, limit=10):

    if not rows:

        return "🔎 Рапорт топилмади."

    shown = rows[:limit]

    text = f"🔎 Топилди: {len(rows)} та рапорт"

    if len(rows) > limit:

        text += f"\n📌 Биринчи {limit} таси кўрсатилмоқда."

    text += "\n\n"

    for row in shown:

        date = row[0] or "—"

        site = row[1] or "Шантиё аниқланмади"

        workers = row[2] if row[2] not in (None, "") else "—"

        report = row[5] or ""

        if len(report) > 250:

            report = report[:250] + "..."
            text += f"📅 {date}\n"

        text += f"📍 {site}\n"

        text += f"👷 {workers}\n"

        text += f"📝 {report}\n\n"
    if not rows:
        return "🔎 Рапорт топилмади."

    shown = rows[:limit]
    t = f"🔎 Топилди: {len(rows)} та рапорт"
    if len(rows) > limit:
        t += f"\nПоказано: биринчи {limit} та"

    t += "\n\n"
    for row in shown:
        date = row[0] or ""
        site = row[1] or "Шантиё кўрсатилмаган"
        workers = row[2] if row[2] not in (None, "") else "—"
        text = str(row[5] or "").replace("\n", " ")
        if len(text) > 220:
            text = text[:220] + "..."

        t += f"📅 {date} | 📍 {site} | 👷 {workers}\n"
        t += f"{text}\n\n"

    return t


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["search_mode"] = False
    await update.message.reply_text(
        "🏗️ Şantiye Rapor Kontrol Botu\n\n"
        "Raporlarınızı gruba normal mesaj olarak gönderebilirsiniz.\n"
        "🔎 Eski raporları tarih, şantiye veya kelime ile arayabilirsiniz.",
        reply_markup=MENU,
    )


async def status(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(status_text(), reply_markup=MENU)


async def search_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["search_mode"] = True
    await update.message.reply_text(
        "🔎 Рапорт қидириш.\n\n"
        "Сана ёзинг: 05.09.2026\n"
        "ёки сана оралиғи: 01.01.2026 - 05.09.2026\n"
        "ёки шантиё: PIRAMIT, BWC, LOT71\n"
        "ёки исталган сўзни ёзинг.",
        reply_markup=ReplyKeyboardMarkup([["❌ Бекор қилиш"]], resize_keyboard=True),
    )


async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.message or not update.message.text:
        return

    text = update.message.text.strip()

    if text == "❌ Бекор қилиш":
        context.user_data["search_mode"] = False
        await update.message.reply_text("Бекор қилинди.", reply_markup=MENU)
        return

    if text == "📊 Rapor Durumu":
        context.user_data["search_mode"] = False
        await status(update, context)
        return

    if text == "📥 Excel":
        context.user_data["search_mode"] = False
        if not ADMIN_CHAT_ID or str(update.effective_chat.id) != str(ADMIN_CHAT_ID):
            await update.message.reply_text("⛔ Bu bölüm sadece yönetici içindir.", reply_markup=MENU)
            return

        ensure_excel()
        with open(DATA_FILE, "rb") as f:
            await update.message.reply_document(f, filename="santiye_raporlari.xlsx")
        return

    if text == "🔎 Rapor Ara":
        await search_start(update, context)
        return

    if context.user_data.get("search_mode"):
        rows = search_history(text)
        await update.message.reply_text(
            format_history_results(rows),
            reply_markup=ReplyKeyboardMarkup(
                [["🔎 Rapor Ara"], ["📊 Rapor Durumu", "📥 Excel"]],
                resize_keyboard=True,
            ),
        )
        return

    site = detect_site(text)
    if not site:
        return

    workers = detect_total_workers(text)
    sender = (
        update.effective_user.full_name
        if update.effective_user
        else "Bilinmiyor"
    )

    save_report(site, sender, workers, text, update.message.message_id)

    if workers is not None:
        await update.message.reply_text(
            f"✅ {site} raporu kaydedildi.\n👷 Toplam: {workers} kişi",
            reply_markup=MENU,
        )
    else:
        await update.message.reply_text(
            f"✅ {site} raporu kaydedildi.\n⚠️ Toplam işçi sayısı raporda bulunamadı.",
            reply_markup=MENU,
        )


async def daily_report(context: ContextTypes.DEFAULT_TYPE):
    if ADMIN_CHAT_ID:
        await context.bot.send_message(
            chat_id=ADMIN_CHAT_ID,
            text=status_text(),
        )


def main():
    if not BOT_TOKEN:
        raise RuntimeError("BOT_TOKEN .env içine girilmelidir.")

    ensure_excel()

    app = Application.builder().token(BOT_TOKEN).build()

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("durum", status))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))

    app.job_queue.run_daily(
        daily_report,
        time=datetime.now(TZ).replace(
            hour=REPORT_HOUR,
            minute=REPORT_MINUTE,
            second=0,
            microsecond=0,
        ).time(),
    )

    print("Bot çalışıyor...")
    app.run_polling()


if __name__ == "__main__":
    main()
